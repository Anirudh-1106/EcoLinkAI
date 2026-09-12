"""
Ingest the real KINFRA company dataset.

Unlike the synthetic files under datasets/csv, this workbook holds data
collected from actual companies in the KINFRA industrial park, Trivandrum.
They are sellers: each has a plant, a waste material and a live listing, but
no recorded buyers and no exchange history, because the companies did not
report who they sell to.

Run this after any database reseed. scripts/database/seed.py only knows about
the synthetic CSVs and will drop these companies, so they have to be put back
separately. Re-running is safe: rows are matched on their natural keys and
updated in place rather than duplicated.
"""

from __future__ import annotations

import os
import sys
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal

import openpyxl

# Add backend to path so we can import models and database
sys.path.append(os.path.join(os.path.dirname(__file__), "backend"))

from app.core.database import SessionLocal
from app.core.security import hash_password
from app.enums.common import QuantityUnit
from app.enums.company import VerificationStatus
from app.enums.exchange import WasteStatus
from app.enums.material import HazardClass, MaterialCategory
from app.models.analytics import Analytics
from app.models.company import Company
from app.models.material import Material
from app.models.plant import Plant
from app.models.user import User, UserRole
from app.models.waste_listing import WasteListing

WORKBOOK = os.path.join(os.path.dirname(__file__), "datasets", "kinfra_data.xlsx")

# The park's real location; plants are nudged apart slightly so they do not
# stack on the exact same pixel on the map.
KINFRA_LAT = 8.554790
KINFRA_LON = 76.858660

# Matches the value used for seeded materials, so carbon savings estimated for
# KINFRA listings are comparable with the rest of the catalogue.
DEFAULT_CARBON_FACTOR = Decimal("1.85")

DEFAULT_PASSWORD = "password123"

CATEGORY_MAP = {
    "rubber": MaterialCategory.RUBBER,
    "wood": MaterialCategory.WOOD,
    "plastic": MaterialCategory.PLASTIC,
    "metal": MaterialCategory.METAL,
    "aluminium": MaterialCategory.METAL,
    "aluminum": MaterialCategory.METAL,
    "paper": MaterialCategory.PAPER,
    "glass": MaterialCategory.GLASS,
    "textile": MaterialCategory.TEXTILE,
    "organic": MaterialCategory.ORGANIC,
    "chemical": MaterialCategory.CHEMICAL,
    "electronic": MaterialCategory.ELECTRONIC,
    "construction": MaterialCategory.CONSTRUCTION,
}

LISTING_STATUS_MAP = {
    "active": WasteStatus.AVAILABLE,
    "available": WasteStatus.AVAILABLE,
    "reserved": WasteStatus.RESERVED,
    "sold": WasteStatus.EXCHANGED,
    "exchanged": WasteStatus.EXCHANGED,
    "expired": WasteStatus.EXPIRED,
}


def _sheet_rows(workbook, name: str) -> list[dict]:
    """Read a worksheet into a list of header-keyed dicts."""
    if name not in workbook.sheetnames:
        raise ValueError(
            f"Sheet {name!r} missing from {WORKBOOK}. "
            f"Found: {workbook.sheetnames}"
        )
    sheet = workbook[name]
    headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    rows = []
    for r in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        if all(v is None for v in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def _as_date(value, fallback: date) -> date:
    """Dates arrive as datetimes or strings depending on how the sheet was saved."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        text = value.strip().split(" ")[0]
        try:
            return date.fromisoformat(text)
        except ValueError:
            return fallback
    return fallback


def _as_decimal(value, fallback: str = "0") -> Decimal:
    try:
        return Decimal(str(value)) if value is not None else Decimal(fallback)
    except Exception:
        return Decimal(fallback)


def ingest_kinfra_data() -> int:
    """Load the workbook into the database. Returns the number of companies handled."""
    workbook = openpyxl.load_workbook(WORKBOOK, data_only=True)

    companies = _sheet_rows(workbook, "companies")
    plants = _sheet_rows(workbook, "plants")
    materials = _sheet_rows(workbook, "materials")
    listings = _sheet_rows(workbook, "waste_listings")

    if not companies:
        raise ValueError(f"No company rows found in {WORKBOOK}; nothing to ingest.")

    db = SessionLocal()
    created = {"companies": 0, "plants": 0, "materials": 0, "listings": 0, "users": 0}
    updated = {"companies": 0, "plants": 0, "listings": 0}

    try:
        # ── Companies ─────────────────────────────────
        company_ids: dict[str, uuid.UUID] = {}
        for idx, row in enumerate(companies):
            sheet_id = str(row["company_id"])
            name = f"KINFRA - {row['company_name']}"
            registration = str(row.get("registration_number") or f"KIN-{sheet_id}")

            company = (
                db.query(Company)
                .filter(Company.registration_number == registration)
                .first()
            )
            verification = (
                VerificationStatus.VERIFIED
                if str(row.get("verification_status", "")).lower() == "verified"
                else VerificationStatus.PENDING
            )

            if company is None:
                company = Company(
                    id=uuid.uuid4(),
                    company_name=name,
                    registration_number=registration,
                    gst_number=f"32{registration[-10:]}1Z5"[:15],
                    license_number=f"LIC-{sheet_id}",
                )
                db.add(company)
                created["companies"] += 1
            else:
                updated["companies"] += 1

            company.company_name = name
            company.industry_type = str(row.get("industry_sector") or "Manufacturing")
            company.verification_status = verification
            company.trust_score = float(row.get("trust_score") or 50.0)
            company.is_active = True

            db.flush()
            company_ids[sheet_id] = company.id

            # A login for each company, otherwise nobody can act on their
            # behalf -- accepting an exchange request requires a user of the
            # selling company.
            email = str(row.get("email") or f"{sheet_id.lower()}@kinfra.in")
            if not db.query(User).filter(User.email == email).first():
                db.add(
                    User(
                        email=email,
                        hashed_password=hash_password(DEFAULT_PASSWORD),
                        full_name=f"{row['company_name']} Manager",
                        role=UserRole.COMPANY,
                        company_id=company.id,
                    )
                )
                created["users"] += 1

        db.flush()

        # ── Plants ────────────────────────────────────
        plant_ids: dict[str, uuid.UUID] = {}
        for idx, row in enumerate(plants):
            sheet_id = str(row["plant_id"])
            company_uuid = company_ids.get(str(row["company_id"]))
            if company_uuid is None:
                continue

            plant_name = str(row["plant_name"])
            plant = (
                db.query(Plant)
                .filter(Plant.company_id == company_uuid, Plant.plant_name == plant_name)
                .first()
            )

            is_new = plant is None
            if plant is None:
                plant = Plant(id=uuid.uuid4(), company_id=company_uuid, plant_name=plant_name)
                created["plants"] += 1
            else:
                updated["plants"] += 1

            # Populate before flushing; several of these columns are NOT NULL.
            plant.plant_type = str(row.get("plant_type") or "Manufacturing Unit")
            plant.address = "KINFRA Industrial Park"
            plant.district = str(row.get("district") or "Trivandrum")
            plant.state = str(row.get("state") or "Kerala")
            plant.country = str(row.get("country") or "India")
            # Spread plants slightly around the park's real coordinates.
            plant.latitude = Decimal(str(round(KINFRA_LAT + (idx % 5) * 0.004, 6)))
            plant.longitude = Decimal(str(round(KINFRA_LON + (idx % 4) * 0.004, 6)))
            plant.is_active = True

            if is_new:
                db.add(plant)
            db.flush()
            if is_new:
                db.add(Analytics(plant_id=plant.id))
            plant_ids[sheet_id] = plant.id

        db.flush()

        # ── Materials ─────────────────────────────────
        # Reused by name where one already exists, so a buyer's requirement for
        # e.g. "Plastic Scrap" matches these sellers rather than sitting beside
        # a duplicate material nobody is linked to.
        material_ids: dict[str, uuid.UUID] = {}
        for row in materials:
            sheet_id = str(row["material_id"])
            name = str(row["material_name"])

            material = db.query(Material).filter(Material.material_name == name).first()
            if material is None:
                category = CATEGORY_MAP.get(
                    str(row.get("material_category") or "").strip().lower(),
                    MaterialCategory.OTHER,
                )
                hazardous = str(row.get("hazardous", "No")).strip().lower() == "yes"
                material = Material(
                    id=uuid.uuid4(),
                    material_name=name,
                    material_category=category,
                    hazard_class=(
                        HazardClass.HAZARDOUS if hazardous else HazardClass.NON_HAZARDOUS
                    ),
                    default_density=Decimal("1.2"),
                    carbon_factor=DEFAULT_CARBON_FACTOR,
                    description=f"{row.get('material_type', 'Waste')} - {name}",
                )
                db.add(material)
                db.flush()
                created["materials"] += 1

            material_ids[sheet_id] = material.id

        db.flush()

        # ── Waste listings (the products these companies sell) ──
        today = date.today()
        for row in listings:
            plant_uuid = plant_ids.get(str(row["plant_id"]))
            material_uuid = material_ids.get(str(row["material_id"]))
            if plant_uuid is None or material_uuid is None:
                continue

            listing = (
                db.query(WasteListing)
                .filter(
                    WasteListing.plant_id == plant_uuid,
                    WasteListing.material_id == material_uuid,
                )
                .first()
            )
            if listing is None:
                listing = WasteListing(
                    id=uuid.uuid4(), plant_id=plant_uuid, material_id=material_uuid
                )
                db.add(listing)
                created["listings"] += 1
            else:
                updated["listings"] += 1

            listing.quantity = _as_decimal(row.get("quantity_kg"))
            try:
                listing.unit = QuantityUnit(str(row.get("unit") or "kg").strip().lower())
            except ValueError:
                listing.unit = QuantityUnit.KG
            listing.purity_percentage = _as_decimal(row.get("purity_percentage"), "0")
            listing.price_per_unit = _as_decimal(row.get("asking_price_per_kg_inr"))
            listing.available_from = _as_date(row.get("availability_date"), today)
            listing.available_until = _as_date(
                row.get("expiry_date"), today + timedelta(days=365)
            )
            listing.status = LISTING_STATUS_MAP.get(
                str(row.get("listing_status") or "").strip().lower(),
                WasteStatus.AVAILABLE,
            )
            listing.description = str(
                row.get("notes") or f"Waste generated at KINFRA facility."
            )

        db.commit()
    finally:
        db.close()

    print("KINFRA ingest complete.")
    print(
        f"  companies : {created['companies']} created, {updated['companies']} updated"
    )
    print(f"  users     : {created['users']} created")
    print(f"  plants    : {created['plants']} created, {updated['plants']} updated")
    print(f"  materials : {created['materials']} created (existing names reused)")
    print(f"  listings  : {created['listings']} created, {updated['listings']} updated")
    return len(companies)


if __name__ == "__main__":
    handled = ingest_kinfra_data()
    if handled == 0:
        raise SystemExit("No KINFRA companies were ingested - check the workbook.")
